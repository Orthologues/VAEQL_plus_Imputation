#########################################################
# Author: Jiawei Zhao
# Email1: jiz@sdu.dk
# Email2: jwz.student.bmc.lu@gmail.com
# Date: 2026-08-02
# Description: AWS Batch boilerplate for Ministral feature-type profiling.
# Development: Mainly written with GPT-5.6 Luna-XHigh on Codex, with Jiawei Zhao's human
# review and revisions. Human review of the generated profile and this boilerplate is required.
#########################################################

"""Step 0 AWS Batch entry point for preliminary feature-type profiling.

The SLM runs inside an approved AWS Batch GPU container. This module only
coordinates bounded schema profiling and writes a pending-human-review manifest;
it is not a clinical metadata authority or a production annotation system.
"""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from pathlib import Path
from typing import Any, Mapping, Sequence

from VAEQL_plus.util.AWS_Batch_interface import (
    SSE_CUSTOMER_KEY_ENV_VAR,
    AWS_Batch_Interface,
)
from VAEQL_plus.util.AWS_S3_interface import AWS_S3_Interface


# Constants
DEFAULT_MODEL_NAME = "mistralai/Ministral-3-8B-Instruct-2512"
ALTERNATIVE_MODEL_NAMES = {
    "mistralai/Ministral-3-14B-Instruct-2512",
    "Qwen/Qwen3-VL-8B-Instruct",
    "google/gemma-4-12B-it",
}
JOB_MODULE = "VAEQL_plus.step0_SLM_metadata_profiling.feature_type_profiling"
ALLOWED_MODEL_TYPES = {
    "continuous",
    "positive_continuous",
    "count",
    "binary",
    "categorical",
    "ordinal",
}
# Document inputs are normalized to bounded Markdown or JSON before prompting.
_TEXT_METADATA_SUFFIXES = {".csv", ".json", ".md", ".tex", ".txt", ".yaml", ".yml"}


# the two primary classes for the workflow
class FeatureTypeProfiler:
    """Minimal local-in-container adapter for Ministral feature-type profiling."""

    @staticmethod
    def build_profiling_prompt(
        schema_summary: Mapping[str, Any],
        metadata_bundle: Sequence[Mapping[str, Any]] = (),
    ) -> str:
        """Build a constrained JSON-only prompt for the Batch-local SLM."""
        return (
            "Profile the tabular schema and approved non-raw metadata below for a "
            "research preprocessing pipeline. "
            "Return JSON only with one object per source feature under `features`. "
            "Use exactly one model_type from continuous, positive_continuous, count, "
            "binary, categorical, or ordinal. Do not invent clinical facts. "
            "Set confidence between 0 and 1 and explain evidence briefly. "
            "Human review is mandatory. Raw row-level data is excluded from this prompt.\n\n"
            + json.dumps(
                {
                    "pandas_metadata_summary": schema_summary,
                    "approved_non_raw_metadata": list(metadata_bundle),
                },
                ensure_ascii=True,
                sort_keys=True,
            )
        )

    @staticmethod
    def parse_profile_response(
        response_text: str,
        source_features: Sequence[str],
    ) -> list[dict[str, Any]]:
        """Validate the SLM JSON response before it reaches preprocessing."""
        start = response_text.find("{")
        end = response_text.rfind("}")
        if start < 0 or end <= start:
            raise ValueError("SLM response did not contain a JSON object")
        try:
            payload = json.loads(response_text[start : end + 1])
        except json.JSONDecodeError as exc:
            raise ValueError("SLM response did not contain valid JSON") from exc

        raw_features = payload.get("features") if isinstance(payload, dict) else None
        if not isinstance(raw_features, list):
            raise ValueError("SLM response must contain a `features` list")

        expected = list(source_features)
        expected_set = set(expected)
        normalized: list[dict[str, Any]] = []
        seen: set[str] = set()
        for item in raw_features:
            if not isinstance(item, dict):
                raise ValueError("Each profiled feature must be a JSON object")
            source_feature = item.get("source_feature")
            model_type = item.get("model_type")
            if source_feature not in expected_set:
                raise ValueError(f"Unknown source feature from SLM: {source_feature!r}")
            if source_feature in seen:
                raise ValueError(f"Duplicate source feature from SLM: {source_feature!r}")
            if model_type not in ALLOWED_MODEL_TYPES:
                raise ValueError(f"Unsupported model_type from SLM: {model_type!r}")
            confidence = item.get("confidence")
            if not isinstance(confidence, (int, float)) or not 0 <= float(confidence) <= 1:
                raise ValueError(f"Invalid confidence for {source_feature!r}")
            seen.add(source_feature)
            normalized.append(
                {
                    "source_feature": source_feature,
                    "canonical_feature": item.get("canonical_feature", source_feature),
                    "model_type": model_type,
                    "num_levels": item.get("num_levels"),
                    "raw_to_canonical": item.get("raw_to_canonical", {}),
                    "missing_value_codes": item.get("missing_value_codes", []),
                    "evidence": item.get("evidence", ""),
                    "confidence": float(confidence),
                }
            )

        missing = expected_set - seen
        if missing:
            raise ValueError(f"SLM response omitted source features: {sorted(missing)!r}")
        return normalized

    def __init__(
        self,
        *,
        model_name: str = DEFAULT_MODEL_NAME,
        max_new_tokens: int = 1024,
    ) -> None:
        if not model_name:
            raise ValueError("model_name must be non-empty")
        if max_new_tokens < 1:
            raise ValueError("max_new_tokens must be greater than zero")
        self.model_name = model_name
        self.max_new_tokens = max_new_tokens

    def _generate_text(self, prompt: str) -> str:
        """Generate text through the loader required by the selected model family."""
        try:
            import torch
        except ModuleNotFoundError as exc:  # pragma: no cover - Batch image dependency
            raise RuntimeError("torch is required in the AWS Batch SLM image") from exc

        messages = [{"role": "user", "content": prompt}]
        if self.model_name.startswith("mistralai/Ministral-3-"):
            try:
                from transformers import Mistral3ForConditionalGeneration, MistralCommonBackend
            except ImportError as exc:  # pragma: no cover - Batch image dependency
                raise RuntimeError(
                    "latest transformers and mistral-common are required for Ministral 3"
                ) from exc
            processor = MistralCommonBackend.from_pretrained(self.model_name)
            model = Mistral3ForConditionalGeneration.from_pretrained(
                self.model_name,
                device_map="auto",
            )
            inputs = processor.apply_chat_template(
                messages,
                return_tensors="pt",
                return_dict=True,
            )
            decoder = processor
        elif self.model_name.startswith(("Qwen/", "google/gemma-4-")):
            try:
                from transformers import AutoModelForMultimodalLM, AutoProcessor
            except ImportError as exc:  # pragma: no cover - Batch image dependency
                raise RuntimeError(
                    "a recent transformers release is required for multimodal SLMs"
                ) from exc
            processor = AutoProcessor.from_pretrained(self.model_name)
            model = AutoModelForMultimodalLM.from_pretrained(
                self.model_name,
                dtype="auto",
                device_map="auto",
            )
            inputs = processor.apply_chat_template(
                messages,
                tokenize=True,
                return_dict=True,
                return_tensors="pt",
                add_generation_prompt=True,
            )
            decoder = processor
        else:
            try:
                from transformers import AutoModelForCausalLM, AutoTokenizer
            except ImportError as exc:  # pragma: no cover - Batch image dependency
                raise RuntimeError("transformers is required in the AWS Batch SLM image") from exc
            decoder = AutoTokenizer.from_pretrained(self.model_name)
            model = AutoModelForCausalLM.from_pretrained(
                self.model_name,
                device_map="auto",
                torch_dtype="auto",
            )
            inputs = decoder(prompt, return_tensors="pt")

        inputs = inputs.to(model.device)
        with torch.inference_mode():
            outputs = model.generate(**inputs, max_new_tokens=self.max_new_tokens)
        input_length = inputs["input_ids"].shape[-1]
        return decoder.decode(outputs[0][input_length:], skip_special_tokens=True)

    def profile(
        self,
        schema_summary: Mapping[str, Any],
        metadata_bundle: Sequence[Mapping[str, Any]] = (),
    ) -> list[dict[str, Any]]:
        """Run Ministral inside the AWS Batch container and validate its JSON output."""
        prompt = self.build_profiling_prompt(schema_summary, metadata_bundle)
        generated = self._generate_text(prompt)
        source_features = [item["source_feature"] for item in schema_summary["features"]]
        return self.parse_profile_response(generated, source_features)


class FeatureTypeProfilingBySLM:
    """Coordinate Step 0 input loading, SLM profiling, and manifest output."""

    @staticmethod
    def _write_attachment(
        data: bytes,
        *,
        output_dir: Path | None,
        file_name: str,
    ) -> Path | None:
        """Write one extracted image when normalized output is requested."""
        if output_dir is None:
            return None
        attachment_dir = output_dir / "assets"
        attachment_dir.mkdir(parents=True, exist_ok=True)
        attachment_path = attachment_dir / file_name
        attachment_path.write_bytes(data)
        return attachment_path

    @staticmethod
    def _write_normalized_file(
        content: str,
        *,
        output_dir: Path | None,
        file_name: str,
    ) -> Path | None:
        """Write one normalized metadata file when an output directory is supplied."""
        if output_dir is None:
            return None
        output_dir.mkdir(parents=True, exist_ok=True)
        normalized_path = output_dir / file_name
        normalized_path.write_text(content, encoding="utf-8")
        return normalized_path

    @staticmethod
    def _markdown_table(rows: Sequence[Sequence[Any]]) -> str:
        """Render bounded tabular values as a Markdown table."""
        if not rows:
            return "_Empty table._"
        width = max(len(row) for row in rows)
        normalized_rows = [list(row) + [""] * (width - len(row)) for row in rows]

        def cell(value: Any) -> str:
            return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")

        header = [cell(value) for value in normalized_rows[0]]
        separator = ["---"] * width
        body = [[cell(value) for value in row] for row in normalized_rows[1:]]
        lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(separator) + " |"]
        lines.extend("| " + " | ".join(row) + " |" for row in body)
        return "\n".join(lines)

    @staticmethod
    def _pdf_to_markdown(
        path: Path,
        *,
        output_dir: Path | None,
        file_prefix: str,
    ) -> tuple[str, list[Path]]:
        """Extract PDF text and embedded images into Markdown attachments."""
        try:
            import fitz
        except ModuleNotFoundError as exc:  # pragma: no cover - optional converter
            raise RuntimeError("pymupdf is required to normalize PDF metadata") from exc

        lines = [f"# {path.stem}"]
        image_paths: list[Path] = []
        with fitz.open(str(path)) as document:
            for page_number, page in enumerate(document, start=1):
                lines.extend((f"\n## Page {page_number}", page.get_text("text").strip()))
                for image_number, image_info in enumerate(page.get_images(full=True)):
                    image = document.extract_image(image_info[0])
                    extension = image.get("ext", "png")
                    image_name = f"{file_prefix}_page{page_number}_{image_number}.{extension}"
                    image_path = FeatureTypeProfilingBySLM._write_attachment(
                        image["image"],
                        output_dir=output_dir,
                        file_name=image_name,
                    )
                    if image_path is None:
                        lines.append(f"_Embedded image {image_name} was not materialized._")
                    else:
                        image_paths.append(image_path)
                        lines.append(f"![{image_name}](assets/{image_name})")
        return "\n".join(lines), image_paths

    @staticmethod
    def _word_to_markdown(
        path: Path,
        *,
        output_dir: Path | None,
        file_prefix: str,
    ) -> tuple[str, list[Path]]:
        """Extract modern or legacy Word content into Markdown attachments."""
        if path.suffix.lower() == ".doc":
            try:
                import textract
            except ModuleNotFoundError as exc:  # pragma: no cover - optional converter
                raise RuntimeError(
                    "textract is required to normalize legacy .doc metadata"
                ) from exc
            content = textract.process(str(path)).decode("utf-8", errors="replace")
            return f"# {path.stem}\n\n{content}", []

        try:
            from docx import Document
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
        except ModuleNotFoundError as exc:  # pragma: no cover - optional converter
            raise RuntimeError("python-docx is required to normalize .docx metadata") from exc

        document = Document(str(path))
        lines = [f"# {path.stem}"]
        lines.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        for table_number, table in enumerate(document.tables, start=1):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            lines.extend((f"\n## Table {table_number}", FeatureTypeProfilingBySLM._markdown_table(rows)))

        image_paths: list[Path] = []
        for image_number, relationship in enumerate(document.part.rels.values()):
            if relationship.reltype != RT.IMAGE:
                continue
            image_part = relationship.target_part
            extension = Path(str(image_part.partname)).suffix.lstrip(".") or "png"
            image_name = f"{file_prefix}_image{image_number}.{extension}"
            image_path = FeatureTypeProfilingBySLM._write_attachment(
                image_part.blob,
                output_dir=output_dir,
                file_name=image_name,
            )
            if image_path is None:
                lines.append(f"_Embedded image {image_name} was not materialized._")
            else:
                image_paths.append(image_path)
                lines.append(f"![{image_name}](assets/{image_name})")
        return "\n".join(lines), image_paths

    @staticmethod
    def _excel_to_markdown(
        path: Path,
        *,
        output_dir: Path | None,
        file_prefix: str,
        max_rows: int,
    ) -> tuple[str, list[Path]]:
        """Extract bounded Excel sheets and embedded images into Markdown."""
        lines = [f"# {path.stem}"]
        image_paths: list[Path] = []
        if path.suffix.lower() == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ModuleNotFoundError as exc:  # pragma: no cover - optional converter
                raise RuntimeError("openpyxl is required to normalize .xlsx metadata") from exc
            workbook = load_workbook(path, data_only=True, read_only=False)
            for sheet_number, worksheet in enumerate(workbook.worksheets, start=1):
                rows = [
                    list(row)
                    for row in worksheet.iter_rows(values_only=True, max_row=max_rows)
                ]
                lines.extend(
                    (
                        f"\n## Sheet {worksheet.title}",
                        FeatureTypeProfilingBySLM._markdown_table(rows),
                    )
                )
                for image_number, image in enumerate(getattr(worksheet, "_images", ())):
                    extension = getattr(image, "format", None) or "png"
                    image_name = f"{file_prefix}_sheet{sheet_number}_{image_number}.{extension}"
                    image_path = FeatureTypeProfilingBySLM._write_attachment(
                        image._data(),
                        output_dir=output_dir,
                        file_name=image_name,
                    )
                    if image_path is not None:
                        image_paths.append(image_path)
                        lines.append(f"![{image_name}](assets/{image_name})")
            workbook.close()
            return "\n".join(lines), image_paths

        try:
            import pandas as pd
        except ModuleNotFoundError as exc:  # pragma: no cover - project dependency
            raise RuntimeError("pandas is required to normalize .xls metadata") from exc
        try:
            sheets = pd.read_excel(path, sheet_name=None, nrows=max_rows)
        except ImportError as exc:  # pragma: no cover - optional converter
            raise RuntimeError("xlrd is required to normalize legacy .xls metadata") from exc
        for sheet_name, frame in sheets.items():
            rows = [list(frame.columns)] + frame.astype(object).where(frame.notna(), "").values.tolist()
            lines.extend((f"\n## Sheet {sheet_name}", FeatureTypeProfilingBySLM._markdown_table(rows)))
        return "\n".join(lines), image_paths

    @staticmethod
    def _dat_to_json(path: Path, *, max_rows: int) -> str:
        """Emit metadata-only JSON for a delimited or opaque `.dat` file."""
        try:
            import pandas as pd
        except ModuleNotFoundError as exc:  # pragma: no cover - project dependency
            raise RuntimeError("pandas is required to normalize .dat metadata") from exc
        try:
            frame = pd.read_csv(path, sep=None, engine="python", nrows=max_rows)
            columns = [
                {
                    "name": str(name),
                    "pandas_dtype": str(frame[name].dtype),
                    "sampled_missing": int(frame[name].isna().sum()),
                    "sampled_unique": int(frame[name].nunique(dropna=True)),
                }
                for name in frame.columns
            ]
            payload = {
                "source_file": path.name,
                "format": "dat",
                "summary_library": "pandas",
                "sampled_rows": int(len(frame)),
                "columns": columns,
            }
        except (OSError, UnicodeDecodeError, ValueError, pd.errors.ParserError):
            payload = {
                "source_file": path.name,
                "format": "dat",
                "size_bytes": path.stat().st_size,
                "parse_status": "opaque_or_unsupported",
            }
        return json.dumps(payload, ensure_ascii=True, sort_keys=True, indent=2)

    @staticmethod
    def summarize_raw_data_with_pandas(
        path: str | Path,
        *,
        max_rows: int = 256,
    ) -> dict[str, Any]:
        """Build a sanitized Pandas summary without sending raw rows to the SLM."""
        if max_rows < 1:
            raise ValueError("max_rows must be greater than zero")
        input_path = Path(path)
        if not input_path.is_file():
            raise FileNotFoundError(input_path)

        try:
            import pandas as pd
        except ModuleNotFoundError as exc:  # pragma: no cover - Batch image dependency
            raise RuntimeError("pandas is required for the Step 0 metadata summary") from exc

        if input_path.suffix.lower() == ".json":
            frame = pd.read_json(input_path)
        else:
            frame = pd.read_csv(input_path, nrows=max_rows)
        frame = frame.head(max_rows)

        columns: list[dict[str, Any]] = []
        for name in frame.columns:
            series = frame[name]
            non_missing = int(series.notna().sum())
            numeric = pd.to_numeric(series, errors="coerce")
            numeric_non_missing = int(numeric.notna().sum())
            column: dict[str, Any] = {
                "source_feature": str(name),
                "pandas_dtype": str(series.dtype),
                "sampled_non_missing": non_missing,
                "sampled_missing": int(series.isna().sum()),
                "sampled_unique": int(series.nunique(dropna=True)),
                "numeric_parse_rate": round(numeric_non_missing / non_missing, 8)
                if non_missing
                else 0.0,
            }
            if numeric_non_missing:
                column["numeric_min"] = float(numeric.min())
                column["numeric_max"] = float(numeric.max())
            columns.append(column)

        return {
            "source_file": input_path.name,
            "summary_library": "pandas",
            "sampled_rows": int(len(frame)),
            "features": columns,
        }

    summarize_tabular_file = summarize_raw_data_with_pandas

    @staticmethod
    def build_metadata_bundle(
        metadata_paths: Sequence[str | Path],
        *,
        source_uris: Sequence[str] = (),
        max_chars_per_file: int = 20_000,
        max_rows: int = 256,
        output_dir: str | Path | None = None,
    ) -> list[dict[str, Any]]:
        """Build the approved non-raw metadata bundle for the SLM prompt."""
        if max_chars_per_file < 1:
            raise ValueError("max_chars_per_file must be greater than zero")
        if max_rows < 1:
            raise ValueError("max_rows must be greater than zero")
        if source_uris and len(source_uris) != len(metadata_paths):
            raise ValueError("source_uris must align one-to-one with metadata_paths")

        normalized_dir = Path(output_dir) if output_dir is not None else None
        bundle: list[dict[str, Any]] = []
        for index, raw_path in enumerate(metadata_paths):
            path = Path(raw_path)
            if not path.is_file():
                raise FileNotFoundError(path)
            item: dict[str, Any] = {
                "source_uri": source_uris[index] if source_uris else None,
                "file_name": path.name,
                "size_bytes": path.stat().st_size,
            }
            suffix = path.suffix.lower()
            image_paths: list[Path] = []
            normalized_file_name: str | None = None
            if suffix == ".pdf":
                content, image_paths = FeatureTypeProfilingBySLM._pdf_to_markdown(
                    path,
                    output_dir=normalized_dir,
                    file_prefix=f"metadata_{index}",
                )
                normalized_file_name = f"{index:03d}_{path.stem}.md"
            elif suffix in {".doc", ".docx"}:
                content, image_paths = FeatureTypeProfilingBySLM._word_to_markdown(
                    path,
                    output_dir=normalized_dir,
                    file_prefix=f"metadata_{index}",
                )
                normalized_file_name = f"{index:03d}_{path.stem}.md"
            elif suffix in {".xls", ".xlsx"}:
                content, image_paths = FeatureTypeProfilingBySLM._excel_to_markdown(
                    path,
                    output_dir=normalized_dir,
                    file_prefix=f"metadata_{index}",
                    max_rows=max_rows,
                )
                normalized_file_name = f"{index:03d}_{path.stem}.md"
            elif suffix == ".dat":
                content = FeatureTypeProfilingBySLM._dat_to_json(path, max_rows=max_rows)
                normalized_file_name = f"{index:03d}_{path.stem}.json"
            elif suffix in _TEXT_METADATA_SUFFIXES:
                content = path.read_text(encoding="utf-8", errors="replace")
            else:
                content = None

            if content is None:
                item["content"] = None
                item["content_note"] = "Binary metadata content is not included in this text prompt."
            else:
                item["content"] = content[:max_chars_per_file]
                item["content_truncated"] = len(content) > max_chars_per_file
                if normalized_file_name is not None:
                    normalized_path = FeatureTypeProfilingBySLM._write_normalized_file(
                        content,
                        output_dir=normalized_dir,
                        file_name=normalized_file_name,
                    )
                    if normalized_path is not None:
                        item["normalized_path"] = str(normalized_path)
                if image_paths:
                    item["image_paths"] = [str(image_path) for image_path in image_paths]
            bundle.append(item)
        return bundle

    @staticmethod
    def build_profile_manifest(
        *,
        source_s3_uri: str | None = None,
        source_reference: str | None = None,
        model_name: str,
        schema_summary: Mapping[str, Any],
        features: Sequence[Mapping[str, Any]],
        metadata_s3_uris: Sequence[str] = (),
    ) -> dict[str, Any]:
        """Build the pending-human-review profile manifest."""
        if not source_s3_uri and not source_reference:
            raise ValueError("source_s3_uri or source_reference must be provided")
        return {
            "schema_version": "0.1",
            "source_s3_uri": source_s3_uri,
            "source_reference": source_reference or source_s3_uri,
            "model_name": model_name,
            "profile_status": "pending_human_review",
            "human_review_required": True,
            "metadata_s3_uris": list(metadata_s3_uris),
            "schema_summary": dict(schema_summary),
            "features": [dict(feature) for feature in features],
        }

    def __init__(
        self,
        *,
        model_name: str = DEFAULT_MODEL_NAME,
        region_name: str | None = None,
        max_rows: int = 256,
        s3_interface: AWS_S3_Interface | None = None,
        profiler: FeatureTypeProfiler | None = None,
    ) -> None:
        if not model_name:
            raise ValueError("model_name must be non-empty")
        if max_rows < 1:
            raise ValueError("max_rows must be greater than zero")
        self.model_name = model_name
        self.region_name = region_name
        self.max_rows = max_rows
        self.s3_interface = s3_interface
        self.profiler = profiler

    def run(
        self,
        input_s3_uri: str | None,
        output_s3_uri: str,
        *,
        input_path: str | Path | None = None,
        metadata_s3_uris: Sequence[str] = (),
        manifest_path: str | Path | None = None,
    ) -> str:
        """Run Step 0 and write its SSE-C-protected review manifest.

        The SSE-C customer key protects S3 object reads and writes; AWS IAM
        credentials still authorize the S3 and Batch API calls.
        """
        if (input_s3_uri is None) == (input_path is None):
            raise ValueError("provide exactly one of input_s3_uri or input_path")

        s3 = self.s3_interface
        if s3 is None:
            sse_customer_key_b64 = os.environ.get(SSE_CUSTOMER_KEY_ENV_VAR)
            if not sse_customer_key_b64:
                raise RuntimeError(
                    f"{SSE_CUSTOMER_KEY_ENV_VAR} must be set in the AWS Batch container environment"
                )
            s3 = AWS_S3_Interface(
                sse_customer_key_b64=sse_customer_key_b64,
                region_name=self.region_name,
            )

        output_bucket, output_key = s3.parse_s3_uri(output_s3_uri)
        with tempfile.TemporaryDirectory(prefix="vaeql_step0_") as temp_dir:
            if input_path is not None:
                local_input_path = Path(input_path)
                if not local_input_path.is_file():
                    raise FileNotFoundError(local_input_path)
            else:
                _, input_key = s3.parse_s3_uri(input_s3_uri or "")
                input_suffix = Path(input_key).suffix.lower() or ".csv"
                local_input_path = s3.download_file(
                    input_s3_uri or "",
                    Path(temp_dir) / f"pds_trial_input{input_suffix}",
                )

            metadata_paths: list[Path] = []
            for index, metadata_uri in enumerate(metadata_s3_uris):
                if input_s3_uri is not None and metadata_uri == input_s3_uri:
                    raise ValueError("metadata_s3_uris must exclude the raw input URI")
                _, metadata_key = s3.parse_s3_uri(metadata_uri)
                metadata_suffix = Path(metadata_key).suffix.lower() or ".metadata"
                metadata_paths.append(
                    s3.download_file(
                        metadata_uri,
                        Path(temp_dir) / f"metadata_{index}{metadata_suffix}",
                    )
                )

            schema_summary = self.summarize_raw_data_with_pandas(
                local_input_path,
                max_rows=self.max_rows,
            )
            metadata_bundle = self.build_metadata_bundle(
                metadata_paths,
                source_uris=metadata_s3_uris,
                max_rows=self.max_rows,
                output_dir=Path(temp_dir) / "normalized_metadata",
            )
            profiler = self.profiler
            if profiler is None:
                profiler = FeatureTypeProfiler(model_name=self.model_name)
            features = profiler.profile(schema_summary, metadata_bundle)
            manifest = self.build_profile_manifest(
                source_s3_uri=input_s3_uri,
                source_reference=(
                    f"mountpoint://{local_input_path.name}"
                    if input_path is not None
                    else input_s3_uri
                ),
                model_name=self.model_name,
                schema_summary=schema_summary,
                features=features,
                metadata_s3_uris=metadata_s3_uris,
            )
            output_uri = s3.write_json_metadata(manifest, output_bucket, output_key)
            if manifest_path is not None:
                local_manifest_path = Path(manifest_path)
                local_manifest_path.parent.mkdir(parents=True, exist_ok=True)
                local_manifest_path.write_text(
                    json.dumps(manifest, ensure_ascii=True, sort_keys=True, indent=2) + "\n",
                    encoding="utf-8",
                )
            return output_uri

    def submit(
        self,
        batch: AWS_Batch_Interface,
        *,
        job_name: str,
        job_queue: str,
        job_definition: str,
        output_s3_uri: str,
        input_s3_uri: str | None = None,
        input_path: str | Path | None = None,
        metadata_s3_uris: Sequence[str] = (),
        sse_customer_key_b64: str | None = None,
        retry_attempts: int | None = 1,
        timeout_seconds: int | None = None,
    ) -> dict[str, Any]:
        """Submit Step 0 to an AWS Batch GPU job definition.

        The supplied job definition must request the GPU resources required by
        the selected SLM image; this method only submits the job.
        """
        if (input_s3_uri is None) == (input_path is None):
            raise ValueError("provide exactly one of input_s3_uri or input_path")
        if input_s3_uri is not None:
            AWS_S3_Interface.parse_s3_uri(input_s3_uri)
        AWS_S3_Interface.parse_s3_uri(output_s3_uri)

        input_arguments = (
            ["--input-uri", input_s3_uri]
            if input_s3_uri
            else ["--input-path", str(input_path)]
        )
        arguments = [*input_arguments, "--output-uri", output_s3_uri]
        for metadata_uri in metadata_s3_uris:
            arguments.extend(("--metadata-uri", metadata_uri))
        arguments.extend(("--model-name", self.model_name))
        return batch.submit_training_step(
            job_name=job_name,
            job_queue=job_queue,
            job_definition=job_definition,
            step_module=JOB_MODULE,
            arguments=arguments,
            sse_customer_key_b64=sse_customer_key_b64,
            retry_attempts=retry_attempts,
            timeout_seconds=timeout_seconds,
        )


def run_feature_type_profiling(
    input_s3_uri: str | None,
    output_s3_uri: str,
    *,
    input_path: str | Path | None = None,
    metadata_s3_uris: Sequence[str] = (),
    model_name: str = DEFAULT_MODEL_NAME,
    region_name: str | None = None,
    max_rows: int = 256,
    manifest_path: str | Path | None = None,
) -> str:
    """Backward-compatible wrapper around :class:`FeatureTypeProfilingBySLM`."""
    return FeatureTypeProfilingBySLM(
        model_name=model_name,
        region_name=region_name,
        max_rows=max_rows,
    ).run(
        input_s3_uri,
        output_s3_uri,
        input_path=input_path,
        metadata_s3_uris=metadata_s3_uris,
        manifest_path=manifest_path,
    )


def submit_feature_type_profiling_job(
    batch: AWS_Batch_Interface,
    *,
    job_name: str,
    job_queue: str,
    job_definition: str,
    output_s3_uri: str,
    input_s3_uri: str | None = None,
    input_path: str | Path | None = None,
    metadata_s3_uris: Sequence[str] = (),
    model_name: str = DEFAULT_MODEL_NAME,
    sse_customer_key_b64: str | None = None,
    retry_attempts: int | None = 1,
    timeout_seconds: int | None = None,
) -> dict[str, Any]:
    """Backward-compatible wrapper around :class:`FeatureTypeProfilingBySLM`."""
    return FeatureTypeProfilingBySLM(model_name=model_name).submit(
        batch,
        job_name=job_name,
        job_queue=job_queue,
        job_definition=job_definition,
        output_s3_uri=output_s3_uri,
        input_s3_uri=input_s3_uri,
        input_path=input_path,
        metadata_s3_uris=metadata_s3_uris,
        sse_customer_key_b64=sse_customer_key_b64,
        retry_attempts=retry_attempts,
        timeout_seconds=timeout_seconds,
    )


def main() -> None:
    """Run the Batch-container entry point for Step 0."""
    parser = argparse.ArgumentParser(description="Profile PDS feature types with Ministral on AWS Batch")
    input_group = parser.add_mutually_exclusive_group(required=True)
    input_group.add_argument("--input-uri", help="SSE-C-protected input S3 URI")
    input_group.add_argument("--input-path", help="Read-only PDS path exposed by an AWS S3 Mountpoint")
    parser.add_argument("--output-uri", required=True, help="SSE-C-protected output manifest S3 URI")
    parser.add_argument("--manifest-path", default=None, help="Optional local copy of the manifest for workflow channels")
    parser.add_argument(
        "--metadata-uri",
        action="append",
        default=[],
        help="Approved non-raw metadata S3 URI; repeat for multiple files",
    )
    parser.add_argument("--model-name", default=DEFAULT_MODEL_NAME)
    parser.add_argument("--region-name", default=None)
    parser.add_argument("--max-rows", type=int, default=256)
    args = parser.parse_args()
    run_feature_type_profiling(
        args.input_uri,
        args.output_uri,
        input_path=args.input_path,
        metadata_s3_uris=args.metadata_uri,
        model_name=args.model_name,
        region_name=args.region_name,
        max_rows=args.max_rows,
        manifest_path=args.manifest_path,
    )


if __name__ == "__main__":
    main()
